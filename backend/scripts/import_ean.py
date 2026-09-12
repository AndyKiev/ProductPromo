"""Bulk-import product types and EAN barcodes from the EAN xlsx export.

The workbook is large (~740k rows), so it is never loaded into memory: a single
streaming pass COPYs the needed columns (article, barcode, product type) into an
UNLOGGED staging table, then set-based INSERT/UPDATE statements populate
`product_type`, `ean` and `product.product_type_id`. Idempotent and re-runnable.

Usage (from backend/):
    poetry run python -m scripts.import_ean
    poetry run python -m scripts.import_ean --xlsx "C:\\...\\EAN_all.xlsx" --truncate
"""
from __future__ import annotations

import argparse
import asyncio
import os
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import asyncpg  # noqa: E402
import openpyxl  # noqa: E402

from backend.config.config import settings  # noqa: E402
from backend.database.init_db import bootstrap  # noqa: E402
from backend.database.db_helper import db_helper  # noqa: E402

DEFAULT_XLSX = r"C:\Users\andre\Desktop\DataSamples\EAN_all_2026 09 11.xlsx"

# 0-based column indexes in the export (header on row 1).
COL_ARTICLE = 0   # артикул  → product.code
COL_EAN = 1       # штрихкод → ean.ean
COL_PRODUCT_TYPE = 7  # тип товара → product_type.code

STAGING_COLUMNS = ["article", "ean", "product_type"]


def _text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _rows(xlsx_path: str):
    workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            yield (_text(row[COL_ARTICLE]), _text(row[COL_EAN]), _text(row[COL_PRODUCT_TYPE]))
    finally:
        workbook.close()


def _dsn() -> str:
    url = settings.db.url
    if not url.startswith("postgresql+asyncpg://"):
        raise SystemExit(f"import_ean only supports postgresql+asyncpg (got '{url}')")
    return url.replace("postgresql+asyncpg://", "postgresql://", 1)


DDL = [
    # product types (names are filled in later by the user)
    """
    INSERT INTO product_type (code)
    SELECT DISTINCT product_type FROM staging_ean
    WHERE product_type IS NOT NULL
    ON CONFLICT (code) DO NOTHING;
    """,
    # barcodes -> product_id (resolved through product.code, never the raw code)
    """
    INSERT INTO ean (product_id, ean)
    SELECT p.id, st.ean
    FROM staging_ean st
    JOIN product p ON p.code = st.article
    WHERE st.ean IS NOT NULL
    ON CONFLICT (ean) DO NOTHING;
    """,
    # product.product_type_id
    """
    UPDATE product p
    SET product_type_id = sub.product_type_id
    FROM (
        SELECT DISTINCT ON (st.article) st.article, pt.id AS product_type_id
        FROM staging_ean st
        JOIN product_type pt ON pt.code = st.product_type
        WHERE st.product_type IS NOT NULL
        ORDER BY st.article
    ) sub
    WHERE p.code = sub.article;
    """,
    "ANALYZE product; ANALYZE product_type; ANALYZE ean;",
]

TRUNCATE = [
    "DELETE FROM ean;",
    "UPDATE product SET product_type_id = NULL;",
    "DELETE FROM product_type;",
]

COUNTS = [
    ("product_type", "SELECT count(*) FROM product_type"),
    ("ean", "SELECT count(*) FROM ean"),
    ("products with ean", "SELECT count(DISTINCT product_id) FROM ean"),
    ("products with product_type", "SELECT count(*) FROM product WHERE product_type_id IS NOT NULL"),
    ("staged rows not matched to a product",
     "SELECT count(*) FROM staging_ean st LEFT JOIN product p ON p.code = st.article WHERE p.id IS NULL"),
]


async def run(xlsx_path: str, truncate: bool) -> None:
    if not os.path.exists(xlsx_path):
        raise SystemExit(f"xlsx not found: {xlsx_path}")

    started = time.perf_counter()
    print("creating tables (idempotent) ...", flush=True)
    await bootstrap()
    await db_helper.dispose()

    conn = await asyncpg.connect(_dsn())
    try:
        await conn.execute("SET synchronous_commit = off;")
        if truncate:
            print("truncating ean / product_type ...", flush=True)
            for statement in TRUNCATE:
                await conn.execute(statement)

        print("staging xlsx (single streaming pass) ...", flush=True)
        await conn.execute("DROP TABLE IF EXISTS staging_ean;")
        await conn.execute(
            "CREATE UNLOGGED TABLE staging_ean (article text, ean text, product_type text);"
        )
        await conn.copy_records_to_table(
            "staging_ean", records=_rows(xlsx_path), columns=STAGING_COLUMNS
        )
        total = await conn.fetchval("SELECT count(*) FROM staging_ean")
        print(f"  staged rows: {total:,}", flush=True)

        for statement in DDL:
            label = " ".join(statement.split())[:60]
            print(f"  > {label} ...", flush=True)
            await conn.execute(statement)

        for label, query in COUNTS:
            print(f"  {label}: {await conn.fetchval(query):,}", flush=True)
    finally:
        try:
            await conn.execute("DROP TABLE IF EXISTS staging_ean;")
        except Exception:
            pass
        await conn.close()

    total_seconds = time.perf_counter() - started
    print(f"done in {total_seconds:0.1f}s", flush=True)


def main() -> None:
    parser = argparse.ArgumentParser(description="Bulk-import EAN barcodes and product types.")
    parser.add_argument("--xlsx", default=DEFAULT_XLSX, help="Path to the source xlsx.")
    parser.add_argument("--truncate", action="store_true", help="Empty ean/product_type before importing.")
    args = parser.parse_args()
    asyncio.run(run(args.xlsx, args.truncate))


if __name__ == "__main__":
    main()
